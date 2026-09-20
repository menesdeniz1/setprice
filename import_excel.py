# -*- coding: utf-8 -*-
import os
import openpyxl
from datetime import datetime
from sqlalchemy.orm import Session

from src.config_loader import ConfigLoader
from src.logger import setup_logger
from src.web_backend.database import engine, Base, SessionLocal
from src.web_backend import models, crud, schemas

def import_from_excel():
    db: Session = SessionLocal()
    
    # 0. Veritabanı tablolarını oluştur
    Base.metadata.create_all(bind=engine)
    
    config_loader = ConfigLoader()
    config = config_loader.load()
    logger = setup_logger(config)
    
    workbook_path = config.get("paths", {}).get("workbook", "./data/pc_setup.xlsx")
    if not os.path.exists(workbook_path):
        logger.error(f"Excel dosyası bulunamadı: {workbook_path}")
        db.close()
        return
        
    logger.info(f"Excel'den göç başlatılıyor: {workbook_path}")
    
    # 1. İlk kullanıcı için parola dış yapılandırmadan sağlanmalıdır.
    default_email = "admin@setprice.com"
    db_user = crud.get_user_by_email(db, default_email)
    if not db_user:
        password = os.environ.get("SETPRICE_INITIAL_ADMIN_PASSWORD", "")
        if len(password) < 16:
            db.close()
            raise RuntimeError("Set SETPRICE_INITIAL_ADMIN_PASSWORD to a unique password of at least 16 characters before creating the initial account.")
        user_in = schemas.UserCreate(email=default_email, password=password)
        db_user = crud.create_user(db, user_in)
        logger.info(f"Varsayılan kullanıcı oluşturuldu: {default_email}")
    else:
        logger.info(f"Kullanıcı zaten mevcut: {default_email}")
        
    # 2. Varsayılan Set oluştur
    set_name = "Mevcut Excel Sistemim"
    db_set = db.query(models.ProductSet).filter(
        models.ProductSet.user_id == db_user.id,
        models.ProductSet.name == set_name
    ).first()
    
    if not db_set:
        set_in = schemas.ProductSetCreate(name=set_name, target_budget=150000.0)
        db_set = crud.create_set(db, set_in=set_in, user_id=db_user.id)
        logger.info(f"Varsayılan set oluşturuldu: '{set_name}'")
    else:
        logger.info(f"Set zaten mevcut: '{set_name}'")
        
    # 3. Excel'den Kaynakları oku
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    sources_sheet_name = config.get("workbook", {}).get("sources_sheet", "Kaynaklar")
    
    if sources_sheet_name not in wb.sheetnames:
        logger.error(f"'{sources_sheet_name}' sayfası bulunamadı!")
        db.close()
        return
        
    ws = wb[sources_sheet_name]
    cols = config.get("workbook", {}).get("sources_columns", {})
    
    header_map = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=1, column=c).value
        if val:
            header_map[val] = c
            
    imported_count = 0
    now = datetime.utcnow()
    
    # Mevcut ürün linklerini çekelim ki mükerrer eklemeyelim
    existing_links = {p.library_product.original_link for p in db_set.products if p.library_product}
    
    for r in range(2, ws.max_row + 1):
        product = ws.cell(row=r, column=header_map.get(cols.get("product"))).value if cols.get("product") in header_map else None
        link = ws.cell(row=r, column=header_map.get(cols.get("link"))).value if cols.get("link") in header_map else None
        
        if not product or not link:
            continue
            
        link_str = str(link).strip()
        if link_str in existing_links:
            continue
            
        category = ws.cell(row=r, column=header_map.get(cols.get("category"))).value if cols.get("category") in header_map else "Diğer"
        seller = ws.cell(row=r, column=header_map.get(cols.get("seller"))).value if cols.get("seller") in header_map else ""
        price_val = ws.cell(row=r, column=header_map.get(cols.get("price"))).value if cols.get("price") in header_map else None
        status_val = ws.cell(row=r, column=header_map.get(cols.get("status"))).value if cols.get("status") in header_map else "OK"
        active_val = ws.cell(row=r, column=header_map.get(cols.get("active"))).value if cols.get("active") in header_map else "Evet"
        locked_val = ws.cell(row=r, column=header_map.get(cols.get("lock"))).value if cols.get("lock") in header_map else "Hayır"
        installment_val = ws.cell(row=r, column=header_map.get(cols.get("installment"))).value if cols.get("installment") in header_map else None
        
        is_active = str(active_val).strip().lower() == "evet"
        is_locked = str(locked_val).strip().lower() == "evet"
        
        # Fiyatı çözümle
        price = None
        if price_val is not None:
            try:
                price = float(price_val)
            except:
                pass
                
        # Kütüphane ürününü getir veya oluştur
        db_lib = crud.get_or_create_library_product(
            db, 
            name=str(product).strip(),
            category=str(category).strip(),
            original_link=link_str,
            user_id=db_user.id
        )
        
        # Kütüphane ürününün canlı durumunu güncelle
        db_lib.current_price = price
        db_lib.current_seller = str(seller).strip() if seller else None
        db_lib.current_installment = str(installment_val).strip() if installment_val else None
        db_lib.status = str(status_val).strip()
        db.commit()
        
        # Set ilişkisini oluştur
        db_product = models.Product(
            set_id=db_set.id,
            library_product_id=db_lib.id,
            locked_price=price if is_locked else None,
            is_active=is_active,
            is_locked=is_locked,
            updated_at=now
        )
        db.add(db_product)
        db.commit()
        db.refresh(db_product)
        
        # İlk fiyat geçmişini ekle
        if price is not None:
            crud.add_price_history(db, db_lib.id, price, db_lib.current_seller)
            
        existing_links.add(link_str)
        imported_count += 1
        
    logger.info(f"Excel göçü tamamlandı. Toplam {imported_count} ürün başarıyla içe aktarıldı.")
    db.close()

if __name__ == "__main__":
    import_from_excel()
