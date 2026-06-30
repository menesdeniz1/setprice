# -*- coding: utf-8 -*-
"""
Doğrulama Test 2 — cloudscraper ile tekrar dene
"""
import sys, io, re, json, time
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import cloudscraper
from bs4 import BeautifulSoup

scraper = cloudscraper.create_scraper(
    browser={'browser': 'chrome', 'platform': 'windows', 'desktop': True}
)

HEADERS = {
    'Accept-Language': 'tr-TR,tr;q=0.9,en-US;q=0.8',
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
}

def check_site(url, site_name, css_selectors=None):
    print(f"\n{'='*60}")
    print(f"{site_name}")
    print(f"{'='*60}")
    print(f"URL: {url[:80]}...")
    try:
        resp = scraper.get(url, headers=HEADERS, timeout=20)
        print(f"HTTP Status: {resp.status_code}")
        print(f"HTML boyutu: {len(resp.text)} karakter")

        if resp.status_code != 200:
            print(f"❌ {resp.status_code} hatası")
            # İlk 500 karakter
            print(f"Yanıt: {resp.text[:500]}")
            return

        soup = BeautifulSoup(resp.text, 'html.parser')

        # JSON-LD
        jsonld_scripts = soup.find_all('script', type='application/ld+json')
        print(f"\nJSON-LD blok: {len(jsonld_scripts)}")
        for i, s in enumerate(jsonld_scripts):
            try:
                data = json.loads(s.string)
                print(f"  #{i+1}: {json.dumps(data, ensure_ascii=False)[:150]}...")
                # offers.price ara
                def find_price(obj, d=0):
                    if d > 5: return None
                    if isinstance(obj, dict):
                        if 'offers' in obj:
                            o = obj['offers']
                            if isinstance(o, dict) and 'price' in o: return o['price']
                            if isinstance(o, list):
                                for x in o:
                                    if isinstance(x, dict) and 'price' in x: return x['price']
                        for v in obj.values():
                            r = find_price(v, d+1)
                            if r: return r
                    if isinstance(obj, list):
                        for x in obj:
                            r = find_price(x, d+1)
                            if r: return r
                    return None
                p = find_price(data)
                if p: print(f"  ✅ FİYAT: {p}")
            except: pass

        # CSS fallback
        if css_selectors:
            for sel in css_selectors:
                els = soup.select(sel)
                if els:
                    print(f"  CSS [{sel}]: ✅ '{els[0].text.strip()[:40]}'")
                else:
                    print(f"  CSS [{sel}]: ❌")

        # Meta
        for attr_name, attr_val in [('itemprop', 'price'), ('property', 'product:price:amount')]:
            m = soup.find('meta', attrs={attr_name: attr_val})
            if m:
                print(f"  Meta {attr_name}='{attr_val}': ✅ {m.get('content', '')[:20]}")

    except Exception as e:
        print(f"❌ HATA: {type(e).__name__}: {e}")

# ============================================================
# TEST A: E-Ticaret siteleri (cloudscraper ile)
# ============================================================
print("TEST A: E-TİCARET SİTELERİ (cloudscraper)")

import openpyxl
wb = openpyxl.load_workbook('data/pc_setup.xlsx', data_only=True)
urls = {}
for s in wb.sheetnames:
    ws = wb[s]
    for r in range(5, ws.max_row + 1):
        link = ws.cell(r, 6).value
        if not link: continue
        for domain in ['amazon.com.tr', 'hepsiburada.com', 'trendyol.com', 'tebilon.com', 'n11.com']:
            if domain in str(link) and domain not in urls:
                urls[domain] = str(link)

css_map = {
    'amazon.com.tr': ['span.a-offscreen', '#corePrice_feature_div .a-offscreen'],
    'hepsiburada.com': ['[data-test-id="price-current-price"]', 'span[data-bind*="price"]'],
    'trendyol.com': ['span.prc-dsc', 'div.product-price-container .prc-dsc'],
    'tebilon.com': ['.product-price', '.price', '[class*="price"]', '[class*="fiyat"]'],
    'n11.com': ['div.newPrice ins', '.unf-p-summary-price .newPrice'],
}

for domain, url in urls.items():
    check_site(url, domain, css_map.get(domain, []))
    time.sleep(3)

# ============================================================
# TEST B: AKAKÇE (cloudscraper ile)
# ============================================================
print("\n\n" + "#" * 60)
print("TEST B: AKAKÇE.COM (cloudscraper)")
print("#" * 60)

search_term = "RTX 5070 Ti 16GB"
akakce_url = f"https://www.akakce.com/arama/?q={search_term.replace(' ', '+')}"
print(f"\nArama: '{search_term}'")
print(f"URL: {akakce_url}")

try:
    resp = scraper.get(akakce_url, headers=HEADERS, timeout=20)
    print(f"HTTP Status: {resp.status_code}")
    print(f"HTML boyutu: {len(resp.text)} karakter")

    if resp.status_code == 200:
        soup = BeautifulSoup(resp.text, 'html.parser')

        # Çeşitli selector'ları dene
        selectors_to_try = [
            ('li.w', 'li.w (ürün listesi)'),
            ('div.p', 'div.p'),
            ('[class*="product"]', 'class contains product'),
            ('[class*="prc"]', 'class contains prc'),
            ('[class*="price"]', 'class contains price'),
            ('a[class*="pr"]', 'a class contains pr'),
            ('.pB', '.pB'),
            ('.pR', '.pR'),
        ]

        for sel, desc in selectors_to_try:
            els = soup.select(sel)
            if els:
                print(f"\n✅ {desc}: {len(els)} eleman")
                for e in els[:3]:
                    text = e.get_text(' ', strip=True)[:80]
                    href = e.get('href', '') or (e.find('a') or {}).get('href', '') if hasattr(e, 'find') else ''
                    print(f"   - {text}")
                    if href:
                        print(f"     link: {href[:60]}")

        # Sayfadaki tüm fiyat benzeri metinleri bul
        print(f"\n--- Fiyat benzeri metinler ---")
        price_pattern = re.compile(r'\d{1,3}(?:[.,]\d{3})*(?:[.,]\d{2})?\s*(?:TL|₺)', re.IGNORECASE)
        text = soup.get_text()
        prices_found = price_pattern.findall(text)
        print(f"Bulunan fiyat desenleri: {len(prices_found)}")
        for pf in prices_found[:10]:
            print(f"   {pf}")

        # Title
        title = soup.find('title')
        print(f"\nSayfa başlığı: {title.text.strip()[:80] if title else 'yok'}")

        # Genel HTML yapısı (body'nin ana container'ı)
        print(f"\n--- HTML Yapı (temizlenmiş, ilk 1500 kar) ---")
        for tag in soup.find_all(['script', 'style', 'noscript']):
            tag.decompose()
        main = soup.find('main') or soup.find('div', id='content') or soup.find('div', class_=re.compile(r'result|search|list'))
        if main:
            print(str(main)[:1500])
        else:
            body = soup.find('body')
            if body:
                print(str(body)[:1500])

    else:
        print(f"❌ Status: {resp.status_code}")
        print(resp.text[:500])

except Exception as e:
    print(f"❌ HATA: {type(e).__name__}: {e}")

print("\n" + "=" * 60)
print("TÜM TESTLER TAMAMLANDI")
print("=" * 60)
