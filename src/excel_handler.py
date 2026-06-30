# -*- coding: utf-8 -*-
import os
import shutil
from datetime import datetime
from typing import List, Optional, Dict
import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from .models import SourceRow, CatalogRow, AlternativeRow

class ProtectedSheetError(Exception):
    pass

class ExcelHandler:
    def __init__(self, config: dict, logger):
        self.config = config
        self.logger = logger
        self.workbook_path = config.get("paths", {}).get("workbook")
        self.wb = None
        
        wb_config = config.get("workbook", {})
        self.writable_sheets = set(wb_config.get("writable_sheets", []))
        self.sources_sheet_name = wb_config.get("sources_sheet", "Kaynaklar")
        self.catalog_sheet_name = wb_config.get("catalog_sheet", "Katalog")
        self.alternatives_sheet_name = wb_config.get("alternatives_sheet", "Muadiller")
        self.history_sheet_name = wb_config.get("history_sheet", "Fiyat_Gecmisi")

    def load(self) -> None:
        """Dosyayı openpyxl ile yükler."""
        if not os.path.exists(self.workbook_path):
            raise FileNotFoundError(f"Excel dosyası bulunamadı: {self.workbook_path}")
            
        try:
            # Sadece okuma için değil, formülleri korumak için data_only=False kullanıyoruz
            self.wb = openpyxl.load_workbook(self.workbook_path, data_only=False)
            self.logger.info(f"Excel dosyası yüklendi: {self.workbook_path}")
        except PermissionError:
            self.logger.error(f"⚠️ {self.workbook_path} dosyası başka bir program tarafından açık. Lütfen kapatıp tekrar deneyin.")
            raise

    def _assert_writable(self, sheet_name: str) -> None:
        if sheet_name not in self.writable_sheets:
            raise ProtectedSheetError(f"'{sheet_name}' sayfasına yazma izni yok!")

    def save(self) -> None:
        """Atomik kaydetme işlemi yapar."""
        if not self.wb:
            return
            
        try:
            self._apply_uniform_font()
        except Exception as e:
            self.logger.warning(f"Font eşitleme hatası: {e}")
            
        tmp_path = self.workbook_path + ".tmp"
        try:
            self.wb.save(tmp_path)
            # Windows'ta os.rename hedef varsa hata verir, o yüzden os.replace kullanıyoruz
            os.replace(tmp_path, self.workbook_path)
            self.logger.info(f"Değişiklikler kaydedildi: {self.workbook_path}")
        except PermissionError:
            self.logger.error(f"⚠️ {self.workbook_path} dosyası açık olduğu için kaydedilemedi!")
            if os.path.exists(tmp_path):
                try: os.remove(tmp_path)
                except: pass
            raise

    def _apply_uniform_font(self) -> None:
        """Korumalı olmayan (botun yazdığı) sayfaların fontunu, diğer sayfalardan biriyle eşitler."""
        # Korumalı olmayan sayfa adları dışındaki ilk sayfayı referans al
        ref_sheet = None
        for s in self.wb.sheetnames:
            if s not in self.writable_sheets:
                ref_sheet = s
                break
                
        if not ref_sheet:
            return
            
        # C5 hücresinin yazı tipini oku (genelde ürün adları veya çeşitler buradadır)
        ref_cell = self.wb[ref_sheet]['C5']
        ref_font = ref_cell.font
        if not ref_font or not ref_font.name:
            return
            
        font_name = ref_font.name
        font_size = ref_font.size or 11
        
        # Writable sayfalardaki tüm hücrelere bu fontu uygula
        for sheet_name in self.writable_sheets:
            if sheet_name not in self.wb.sheetnames:
                continue
            ws = self.wb[sheet_name]
            for r_idx, row in enumerate(ws.iter_rows(), 1):
                for cell in row:
                    # 1. satır (başlıklar) kalın olsun, diğerleri normal
                    is_header = (r_idx == 1)
                    cell.font = Font(
                        name=font_name,
                        size=font_size,
                        bold=is_header,
                        italic=False,
                        color=ref_font.color
                    )

    def create_backup(self) -> Optional[str]:
        """Otomatik yedek alır."""
        if not os.path.exists(self.workbook_path):
            return None
            
        backups_dir = self.config.get("paths", {}).get("backups_dir", "./data/backups")
        if not os.path.exists(backups_dir):
            os.makedirs(backups_dir)
            
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = os.path.basename(self.workbook_path)
        name, ext = os.path.splitext(filename)
        backup_path = os.path.join(backups_dir, f"{name}_{timestamp}{ext}")
        
        shutil.copy2(self.workbook_path, backup_path)
        self.logger.info(f"Yedek alındı: {backup_path}")
        
        # Eski yedekleri temizle
        max_count = self.config.get("backups", {}).get("max_count", 10)
        backups = []
        for f in os.listdir(backups_dir):
            if f.endswith(ext) and f.startswith(f"{name}_"):
                backups.append(os.path.join(backups_dir, f))
                
        if len(backups) > max_count:
            backups.sort(key=os.path.getmtime)
            for old_backup in backups[:-max_count]:
                try:
                    os.remove(old_backup)
                    self.logger.debug(f"Eski yedek silindi: {old_backup}")
                except Exception as e:
                    self.logger.warning(f"Eski yedek silinemedi: {e}")
                    
        return backup_path

    def create_sheet_if_missing(self, name: str, columns: list) -> None:
        """Belirtilen sayfa yoksa oluşturur ve başlıkları yazar."""
        if name not in self.wb.sheetnames:
            self._assert_writable(name)
            ws = self.wb.create_sheet(title=name)
            for i, col_name in enumerate(columns, 1):
                ws.cell(row=1, column=i, value=col_name)
            self.logger.info(f"Sayfa oluşturuldu: {name}")

    def read_sources(self) -> List[SourceRow]:
        """Kaynaklar sayfasından linkleri okur."""
        if self.sources_sheet_name not in self.wb.sheetnames:
            return []
            
        ws = self.wb[self.sources_sheet_name]
        cols = self.config.get("workbook", {}).get("sources_columns", {})
        
        # Başlık sırasını bul
        header_map = {}
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=c).value
            if val:
                header_map[val] = c
                
        rows = []
        for r in range(2, ws.max_row + 1):
            # Boş satırı atla
            product = ws.cell(row=r, column=header_map.get(cols.get("product"))).value if cols.get("product") in header_map else None
            link = ws.cell(row=r, column=header_map.get(cols.get("link"))).value if cols.get("link") in header_map else None
            
            if not product and not link:
                continue
                
            active_val = ws.cell(row=r, column=header_map.get(cols.get("active"))).value if cols.get("active") in header_map else "Evet"
            locked_val = ws.cell(row=r, column=header_map.get(cols.get("lock"))).value if cols.get("lock") in header_map else "Hayır"
            price_val = ws.cell(row=r, column=header_map.get(cols.get("price"))).value if cols.get("price") in header_map else None
            status_val = ws.cell(row=r, column=header_map.get(cols.get("status"))).value if cols.get("status") in header_map else ""
            
            row_data = SourceRow(
                product=str(product).strip() if product else "",
                category=str(ws.cell(row=r, column=header_map.get(cols.get("category"))).value).strip() if cols.get("category") in header_map else "",
                seller=str(ws.cell(row=r, column=header_map.get(cols.get("seller"))).value).strip() if cols.get("seller") in header_map else "",
                link=str(link).strip() if link else "",
                price=price_val,
                installment=str(ws.cell(row=r, column=header_map.get(cols.get("installment"))).value) if cols.get("installment") in header_map else None,
                status=str(status_val) if status_val else "",
                active=str(active_val).strip().lower() == "evet" if active_val else True,
                locked=str(locked_val).strip().lower() == "evet" if locked_val else False,
                row_number=r
            )
            rows.append(row_data)
            
        return rows

    def write_source_price(self, row_num: int, price: Optional[float], status: str, timestamp: datetime, installment: Optional[str] = None) -> None:
        """Kaynaklar sayfasına fiyat yazar."""
        self._assert_writable(self.sources_sheet_name)
        ws = self.wb[self.sources_sheet_name]
        cols = self.config.get("workbook", {}).get("sources_columns", {})
        
        header_map = {}
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=c).value
            if val: header_map[val] = c
            
        if cols.get("price") in header_map and price is not None:
            ws.cell(row=row_num, column=header_map[cols["price"]], value=price)
            
        if cols.get("installment") in header_map and installment is not None:
            ws.cell(row=row_num, column=header_map[cols["installment"]], value=installment)
            
        if cols.get("status") in header_map:
            ws.cell(row=row_num, column=header_map[cols["status"]], value=status)
            
        if cols.get("updated") in header_map:
            ws.cell(row=row_num, column=header_map[cols["updated"]], value=timestamp.strftime("%Y-%m-%d %H:%M:%S"))

    def read_catalog(self) -> List[CatalogRow]:
        """Katalog'u okur."""
        if self.catalog_sheet_name not in self.wb.sheetnames:
            return []
            
        ws = self.wb[self.catalog_sheet_name]
        cols = self.config.get("workbook", {}).get("catalog_columns", {})
        
        header_map = {}
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=c).value
            if val: header_map[val] = c
            
        rows = []
        for r in range(2, ws.max_row + 1):
            product = ws.cell(row=r, column=header_map.get(cols.get("product"))).value if cols.get("product") in header_map else None
            if not product: continue
            
            row_data = CatalogRow(
                product=str(product).strip(),
                category=str(ws.cell(row=r, column=header_map.get(cols.get("category"))).value) if cols.get("category") in header_map else "",
                best_price=ws.cell(row=r, column=header_map.get(cols.get("best_price"))).value if cols.get("best_price") in header_map else None,
                best_seller=str(ws.cell(row=r, column=header_map.get(cols.get("best_seller"))).value) if cols.get("best_seller") in header_map else "",
                best_link=str(ws.cell(row=r, column=header_map.get(cols.get("best_link"))).value) if cols.get("best_link") in header_map else "",
                search_term=str(ws.cell(row=r, column=header_map.get(cols.get("search_term"))).value) if cols.get("search_term") in header_map else "",
                trend=str(ws.cell(row=r, column=header_map.get(cols.get("trend"))).value) if cols.get("trend") in header_map else "▬"
            )
            rows.append(row_data)
        return rows

    def write_catalog_row(self, product: str, category: str, price: Optional[float], 
                          seller: str, link: str, timestamp: datetime, search_term: str = "", trend: str = "▬") -> None:
        """Katalog'a en ucuz fiyatı yazar. Ürün yoksa ekler."""
        self._assert_writable(self.catalog_sheet_name)
        ws = self.wb[self.catalog_sheet_name]
        cols = self.config.get("workbook", {}).get("catalog_columns", {})
        
        header_map = {}
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=c).value
            if val: header_map[val] = c
            
        target_row = None
        for r in range(2, ws.max_row + 1):
            p = ws.cell(row=r, column=header_map.get(cols.get("product"))).value
            if p and str(p).strip().lower() == product.lower():
                target_row = r
                break
                
        if target_row is None:
            target_row = ws.max_row + 1
            if cols.get("product") in header_map:
                ws.cell(row=target_row, column=header_map[cols["product"]], value=product)
            if cols.get("category") in header_map:
                ws.cell(row=target_row, column=header_map[cols["category"]], value=category)
            if cols.get("search_term") in header_map and search_term:
                ws.cell(row=target_row, column=header_map[cols["search_term"]], value=search_term)
                
        if cols.get("best_price") in header_map and price is not None:
            ws.cell(row=target_row, column=header_map[cols["best_price"]], value=price)
        if cols.get("best_seller") in header_map:
            ws.cell(row=target_row, column=header_map[cols["best_seller"]], value=seller)
        if cols.get("best_link") in header_map:
            ws.cell(row=target_row, column=header_map[cols["best_link"]], value=link)
        if cols.get("trend") in header_map:
            ws.cell(row=target_row, column=header_map[cols["trend"]], value=trend)
        if cols.get("updated") in header_map:
            ws.cell(row=target_row, column=header_map[cols["updated"]], value=timestamp.strftime("%Y-%m-%d %H:%M:%S"))

    def write_alternatives(self, alternatives: List[AlternativeRow]) -> None:
        """Muadiller sayfasını sıfırlar ve yeniden yazar."""
        if not alternatives: return
        
        self._assert_writable(self.alternatives_sheet_name)
        ws = self.wb[self.alternatives_sheet_name]
        cols = self.config.get("workbook", {}).get("alternatives_columns", {})
        
        # Başlık hariç her şeyi sil
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
            
        header_map = {}
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=c).value
            if val: header_map[val] = c
            
        for r_idx, alt in enumerate(alternatives, 2):
            if cols.get("product") in header_map: ws.cell(row=r_idx, column=header_map[cols["product"]], value=alt.product)
            if cols.get("search_term") in header_map: ws.cell(row=r_idx, column=header_map[cols["search_term"]], value=alt.search_term)
            if cols.get("rank") in header_map: ws.cell(row=r_idx, column=header_map[cols["rank"]], value=alt.rank)
            if cols.get("alt_product") in header_map: ws.cell(row=r_idx, column=header_map[cols["alt_product"]], value=alt.alt_product)
            if cols.get("alt_price") in header_map: ws.cell(row=r_idx, column=header_map[cols["alt_price"]], value=alt.alt_price)
            if cols.get("alt_seller") in header_map: ws.cell(row=r_idx, column=header_map[cols["alt_seller"]], value=alt.alt_seller)
            if cols.get("alt_link") in header_map: ws.cell(row=r_idx, column=header_map[cols["alt_link"]], value=alt.alt_link)
            if cols.get("updated") in header_map: ws.cell(row=r_idx, column=header_map[cols["updated"]], value=alt.updated.strftime("%Y-%m-%d %H:%M:%S"))

    def append_history(self, product: str, seller: str, price: float, timestamp: datetime) -> None:
        self._assert_writable(self.history_sheet_name)
        ws = self.wb[self.history_sheet_name]
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=product)
        ws.cell(row=r, column=2, value=seller)
        ws.cell(row=r, column=3, value=price)
        ws.cell(row=r, column=4, value=timestamp.strftime("%Y-%m-%d %H:%M:%S"))

    def apply_setup_formulas(self, dry_run: bool = True) -> List[dict]:
        """Setup sayfalarındaki Fiyat ve Link hücrelerini formülle değiştirir."""
        changes = []
        
        for sheet_name in self.wb.sheetnames:
            if sheet_name in self.writable_sheets:
                continue
                
            ws = self.wb[sheet_name]
            
            # Sütun harflerini belirle (A=1, B=2, C=3 (Ürün), D=4 (Fiyat), F=6 (Link))
            # pc_setup.xlsx analizi göstermişti: Ürün=C, Fiyat=D, Link=F
            product_col_letter = 'C'
            price_col_letter = 'D'
            link_col_letter = 'F'
            
            start_row = self.config.get("migration", {}).get("data_start_row", 5)
            
            for r in range(start_row, ws.max_row + 1):
                product_val = ws.cell(row=r, column=3).value
                if not product_val: continue
                
                # Toplam satırları atla
                skip_patterns = self.config.get("migration", {}).get("skip_product_patterns", [])
                if any(p.lower() in str(product_val).lower() for p in skip_patterns):
                    continue
                    
                price_formula = f'=INDEX({self.catalog_sheet_name}!$C:$C, MATCH(${product_col_letter}{r}, {self.catalog_sheet_name}!$A:$A, 0))'
                link_formula = f'=INDEX({self.catalog_sheet_name}!$E:$E, MATCH(${product_col_letter}{r}, {self.catalog_sheet_name}!$A:$A, 0))'
                trend_formula = f'=IFERROR(INDEX({self.catalog_sheet_name}!$H:$H, MATCH(${product_col_letter}{r}, {self.catalog_sheet_name}!$A:$A, 0)), "▬")'
                
                old_price = ws.cell(row=r, column=4).value
                old_link = ws.cell(row=r, column=6).value
                
                changes.append({
                    "sheet": sheet_name,
                    "row": r,
                    "product": product_val,
                    "old_price": old_price,
                    "new_price_formula": price_formula,
                    "old_link": old_link,
                    "new_link_formula": link_formula
                })
                
                if not dry_run:
                    # Uyarı: Bu işlem setup sayfalarına yazar!
                    ws.cell(row=3, column=7, value="Trend")
                    ws.cell(row=r, column=4, value=price_formula)
                    ws.cell(row=r, column=6, value=link_formula)
                    ws.cell(row=r, column=7, value=trend_formula)
                    
        return changes

    def add_data_validation_to_sources(self) -> None:
        """Kaynaklar sayfasındaki Ürün sütununa, Katalog sayfasındaki Ürün sütununu veri doğrulama olarak ekler."""
        if self.sources_sheet_name not in self.wb.sheetnames or self.catalog_sheet_name not in self.wb.sheetnames:
            return
            
        self._assert_writable(self.sources_sheet_name)
        ws_sources = self.wb[self.sources_sheet_name]
        ws_catalog = self.wb[self.catalog_sheet_name]
        
        # Katalog sayfasındaki ürünlerin aralığı (A2:A{max_row})
        max_cat_row = ws_catalog.max_row
        if max_cat_row < 2:
            return
            
        # Formül: Katalog sayfasındaki A sütunu (Ürün sütunu)
        formula = f"={self.catalog_sheet_name}!$A$2:$A${max_cat_row}"
        
        dv = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="Geçersiz Ürün",
            error="Lütfen Katalog sayfasında tanımlı geçerli bir ürün seçin."
        )
        
        # Kaynaklar sayfasındaki Ürün sütununu bul
        cols = self.config.get("workbook", {}).get("sources_columns", {})
        header_map = {}
        for c in range(1, ws_sources.max_column + 1):
            val = ws_sources.cell(row=1, column=c).value
            if val: header_map[val] = c
            
        prod_col_idx = header_map.get(cols.get("product"))
        if not prod_col_idx:
            return
            
        col_letter = get_column_letter(prod_col_idx)
        
        # A2:A{max_row} şeklinde hücre aralığı ekle
        dv.add(f"{col_letter}2:{col_letter}{ws_sources.max_row + 100}")
        
        # Sayfadaki mevcut validation'ları temizleyip bunu ekleyelim
        ws_sources.data_validations.dataValidation = [dv]
        self.logger.info("Kaynaklar sayfasına veri doğrulama (açılır liste) uygulandı.")

    def sort_sheet_by_category(self, sheet_name: str) -> None:
        """Belirtilen sayfayı kategoriye ve ürün adına göre sıralar."""
        self._assert_writable(sheet_name)
        if sheet_name not in self.wb.sheetnames:
            return
            
        ws = self.wb[sheet_name]
        if ws.max_row <= 2:
            return
            
        # Sütun haritasını oku
        cols_key = "sources_columns" if sheet_name == self.sources_sheet_name else "catalog_columns"
        cols = self.config.get("workbook", {}).get(cols_key, {})
        
        header_map = {}
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=1, column=c).value
            if val:
                header_map[val] = c
                
        cat_col = header_map.get(cols.get("category"))
        prod_col = header_map.get(cols.get("product"))
        
        if not cat_col or not prod_col:
            return
            
        # Satır verilerini topla
        rows_data = []
        for r in range(2, ws.max_row + 1):
            row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            if all(v is None for v in row_vals):
                continue
            rows_data.append(row_vals)
            
        # Kategori ve Ürün adına göre sırala
        def sort_key(row_vals):
            cat = str(row_vals[cat_col - 1] or "").strip().lower()
            prod = str(row_vals[prod_col - 1] or "").strip().lower()
            return (cat, prod)
            
        rows_data.sort(key=sort_key)
        
        # Sayfadaki eski verileri temizle
        for r in range(2, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                ws.cell(row=r, column=c).value = None
                
        # Sıralanmış verileri geri yaz
        for r_idx, row_vals in enumerate(rows_data, 2):
            for c_idx, val in enumerate(row_vals, 1):
                ws.cell(row=r_idx, column=c_idx).value = val
                
        self.logger.info(f"'{sheet_name}' sayfası kategoriye göre sıralandı.")

