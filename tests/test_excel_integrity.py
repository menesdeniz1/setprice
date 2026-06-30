import pytest
import os
import shutil
import openpyxl
from src.excel_handler import ExcelHandler, ProtectedSheetError
from unittest.mock import MagicMock

@pytest.fixture
def config():
    return {
        "paths": {"workbook": "test_data.xlsx"},
        "workbook": {
            "writable_sheets": ["Kaynaklar", "Katalog", "Muadiller"],
            "sources_sheet": "Kaynaklar",
            "catalog_sheet": "Katalog"
        }
    }

@pytest.fixture
def test_excel(config):
    # Geçici bir excel oluştur
    wb = openpyxl.Workbook()
    
    # Setup sayfası (korumalı)
    ws_setup = wb.active
    ws_setup.title = "139k"
    ws_setup.cell(row=5, column=3, value="Test Ürün")
    
    wb.save("test_data.xlsx")
    
    handler = ExcelHandler(config, MagicMock())
    handler.load()
    yield handler
    
    if os.path.exists("test_data.xlsx"):
        os.remove("test_data.xlsx")
    if os.path.exists("test_data.xlsx.tmp"):
        os.remove("test_data.xlsx.tmp")

def test_protected_sheet_write(test_excel):
    # Korumalı sayfaya yazmaya çalışınca hata vermeli
    test_excel.wb["139k"].cell(row=5, column=4, value=999)
    # _assert_writable exception fırlatmalı
    with pytest.raises(ProtectedSheetError):
        test_excel._assert_writable("139k")

def test_create_sheets(test_excel):
    test_excel.create_sheet_if_missing("Kaynaklar", ["Ürün", "Fiyat"])
    assert "Kaynaklar" in test_excel.wb.sheetnames
    assert test_excel.wb["Kaynaklar"].cell(row=1, column=1).value == "Ürün"

def test_atomic_save(test_excel):
    test_excel.create_sheet_if_missing("Kaynaklar", ["Ürün"])
    test_excel.save()
    
    # tmp dosyası silinmiş ve asıl dosya var olmalı
    assert not os.path.exists("test_data.xlsx.tmp")
    assert os.path.exists("test_data.xlsx")
